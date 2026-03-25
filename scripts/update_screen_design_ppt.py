from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


ROOT = Path(r"C:\Users\8316-31\Desktop\기타\개발산출물")
SOURCE = ROOT / "2520110181_김현규_화면설계서_작업사본.pptx"
TARGET = ROOT / "2520110181_김현규_화면설계서_수정본.pptx"
SCREENSHOT_DIR = Path(r"C:\Users\8316-31\SpringWorks\wemeet\artifacts\ppt_screenshots")


SLIDE_UPDATES = {
    30: {
        "title": "게스트 추천 입력",
        "path": "메인페이지 비로그인 > 게스트 추천 입력",
        "actor": "사용자",
        "screen_id": "GUEST_PLAN_001",
        "image": SCREENSHOT_DIR / "guest-plan.png",
    },
    31: {
        "title": "비회원 추천 결과",
        "path": "메인페이지 비로그인 > 추천 결과",
        "actor": "사용자",
        "screen_id": "GUEST_RESULT_001",
        "image": SCREENSHOT_DIR / "guest-results.png",
    },
    32: {
        "title": "친구 관리",
        "path": "메인페이지 로그인 > 친구 관리",
        "actor": "사용자",
        "screen_id": "FRIEND_001",
        "image": SCREENSHOT_DIR / "friends.png",
    },
    33: {
        "title": "검색 이력",
        "path": "메인페이지 로그인 > 검색 이력",
        "actor": "사용자",
        "screen_id": "HISTORY_001",
        "image": SCREENSHOT_DIR / "history.png",
    },
    34: {
        "title": "프로필",
        "path": "메인페이지 로그인 > 프로필",
        "actor": "사용자",
        "screen_id": "PROFILE_001",
        "image": SCREENSHOT_DIR / "profile.png",
    },
    35: {
        "title": "모임 생성",
        "path": "메인페이지 로그인 > 모임 생성",
        "actor": "사용자",
        "screen_id": "MEETING_001",
        "image": SCREENSHOT_DIR / "meeting-form.png",
    },
    36: {
        "title": "추천 결과",
        "path": "메인페이지 로그인 > 모임 생성 > 추천 결과",
        "actor": "사용자",
        "screen_id": "RESULT_001",
        "image": SCREENSHOT_DIR / "member-results.png",
    },
}


def delete_shape(shape):
    shape._element.getparent().remove(shape._element)


def delete_slide(prs: Presentation, slide_index_zero_based: int):
    slide_id_list = prs.slides._sldIdLst
    slides = list(slide_id_list)
    slide_id = slides[slide_index_zero_based]
    rel_id = slide_id.rId
    prs.part.drop_rel(rel_id)
    slide_id_list.remove(slide_id)


def replace_picture(slide, image_path: Path):
    picture = None
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            picture = shape
            break
    if picture is None:
        raise RuntimeError("Picture placeholder not found")

    left = picture.left
    top = picture.top
    width = picture.width
    height = picture.height
    delete_shape(picture)
    slide.shapes.add_picture(str(image_path), left, top, width=width, height=height)


def keep_only_core_shapes(slide):
    placeholders = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER]
    picture = next((shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE), None)
    keep_elements = {shape._element for shape in placeholders[:4]}
    if picture is not None:
        keep_elements.add(picture._element)

    for shape in list(slide.shapes):
        if shape._element not in keep_elements:
            delete_shape(shape)


def set_placeholder_texts(slide, title: str, path_text: str, actor: str, screen_id: str):
    placeholders = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER]
    texts = [title, path_text, actor, screen_id]
    for shape, text in zip(placeholders[:4], texts):
        shape.text = text


def main():
    prs = Presentation(str(SOURCE))

    for slide_number, spec in SLIDE_UPDATES.items():
        slide = prs.slides[slide_number - 1]
        keep_only_core_shapes(slide)
        set_placeholder_texts(
            slide,
            spec["title"],
            spec["path"],
            spec["actor"],
            spec["screen_id"],
        )
        replace_picture(slide, spec["image"])

    for index in range(len(prs.slides), 36, -1):
        delete_slide(prs, index - 1)

    prs.save(str(TARGET))
    print(TARGET)


if __name__ == "__main__":
    main()
