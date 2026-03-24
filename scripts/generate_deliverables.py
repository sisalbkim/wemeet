from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OUTPUT_PATH = Path("artifacts/wemeet-development-deliverables.xlsx")

TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
SECTION_FILL = PatternFill("solid", fgColor="EAF2F8")
SUBSECTION_FILL = PatternFill("solid", fgColor="F5F9FC")
THIN_BORDER = Border(
    left=Side(style="thin", color="B7C3D0"),
    right=Side(style="thin", color="B7C3D0"),
    top=Side(style="thin", color="B7C3D0"),
    bottom=Side(style="thin", color="B7C3D0"),
)


MENU_HEADERS = ["구분", "1depth", "2depth", "3depth", "4depth", "5depth", "Tab/Page", "Page/기능구분", "비고"]
MENU_ROWS = [
    ["Page", "랜딩", "메인 랜딩", "-", "-", "-", "Page", "Program", "카테고리 목록 조회: /api/categories"],
    ["Page", "랜딩", "게스트 추천 입력", "-", "-", "-", "Page", "Program", "카테고리 목록 조회: /api/categories"],
    ["Page", "랜딩", "게스트 추천 결과", "-", "-", "-", "Page", "Program", "추천 결과 조회: /api/recommendations"],
    ["Page", "로그인/회원가입", "로그인", "-", "-", "-", "Page", "Program", "로그인 처리: /api/auth/login"],
    ["Page", "로그인/회원가입", "회원가입", "-", "-", "-", "Page", "Program", "회원가입 처리: /api/auth/signup"],
    ["Page", "메인", "홈", "-", "-", "-", "Page", "Program", "내 정보/친구/모임 조회 사용"],
    ["Page", "메인", "친구 관리", "친구 목록 조회", "-", "-", "Page", "Program", "친구 목록 조회: /api/friends GET"],
    ["Page", "메인", "친구 관리", "친구 코드로 친구 추가", "-", "-", "Page", "Program", "친구 추가: /api/friends POST"],
    ["Page", "메인", "검색 이력", "카테고리 필터 조회", "-", "-", "Tab", "Program", "이력/카테고리 조회: /api/history, /api/categories"],
    ["Page", "메인", "검색 이력", "키워드 검색", "-", "-", "Tab", "Program", "검색 이력 조회: /api/history"],
    ["Page", "메인", "프로필", "프로필 조회", "-", "-", "Page", "Program", "내 정보 조회: /api/me"],
    ["Page", "메인", "프로필", "기본 출발지 수정", "-", "-", "Page", "Program", "주소 수정: /api/me/address"],
    ["Page", "메인", "모임 생성", "참가자 선택", "-", "-", "Page", "Program", "친구/카테고리 조회 사용"],
    ["Page", "메인", "모임 생성", "추천 미리보기", "-", "-", "Page", "Program", "추천 결과 조회: /api/recommendations"],
    ["Page", "메인", "추천 결과", "중간지점 계산", "-", "-", "Page", "Program", "추천 결과 조회: /api/recommendations"],
    ["Page", "메인", "추천 결과", "장소 후보 조회", "-", "-", "Page", "Program", "추천 결과 조회: /api/recommendations"],
    ["Page", "메인", "추천 결과", "지도 포인트 표시", "-", "-", "Page", "Program", "추천 결과 조회: /api/recommendations"],
]

PROGRAM_HEADERS = ["NO", "업무영역 Lev1", "Lev2", "Lev3", "Lev4", "Lev5", "프로그램명", "DB Transaction", "프로그램ID", "작업자", "구분"]
PROGRAM_ROWS = [
    [1, "랜딩", "랜딩 페이지", "게스트 추천 결과", "-", "-", "랜딩/게스트 추천", "R", "LANDING_001", "본인", "사용자"],
    [2, "로그인/회원가입", "로그인", "-", "-", "-", "로그인 화면/처리", "R", "LOGIN_001", "본인", "사용자"],
    [3, "로그인/회원가입", "회원가입", "-", "-", "-", "회원가입 화면/처리", "C", "SIGNUP_001", "본인", "사용자"],
    [4, "로그인/회원가입", "로그아웃", "-", "-", "-", "로그아웃", "-", "LOGOUT_001", "본인", "사용자"],
    [5, "메인", "홈", "-", "-", "-", "홈 대시보드", "R", "HOME_001", "본인", "사용자"],
    [6, "메인", "친구 관리", "친구 목록 조회", "-", "-", "친구 목록", "R", "FRIEND_001", "본인", "사용자"],
    [7, "메인", "친구 관리", "친구 코드로 친구 추가", "-", "-", "친구 추가", "C", "FRIEND_002", "본인", "사용자"],
    [8, "메인", "검색 이력", "이력 조회/필터링", "-", "-", "검색 이력", "R", "HISTORY_001", "본인", "사용자"],
    [9, "메인", "프로필", "프로필 조회", "-", "-", "프로필 조회", "R", "PROFILE_001", "본인", "사용자"],
    [10, "메인", "프로필", "기본 출발지 수정", "-", "-", "기본 주소 수정", "U", "PROFILE_002", "본인", "사용자"],
    [11, "메인", "모임 생성", "모임 작성 화면", "-", "-", "모임 생성 화면", "R", "MEETING_001", "본인", "사용자"],
    [12, "메인", "모임 생성", "추천 미리보기", "-", "-", "추천 미리보기", "R", "MEETING_002", "본인", "사용자"],
    [13, "메인", "모임", "모임 생성", "-", "-", "모임 저장", "C", "MEETING_003", "본인", "사용자"],
    [14, "메인", "추천 결과", "중간지점 계산", "장소 후보 산출", "-", "추천 결과 조회", "R", "RECOMMEND_001", "본인", "사용자"],
    [15, "API", "인증 API", "회원가입", "-", "-", "/api/auth/signup", "C", "API_AUTH_001", "본인", "API"],
    [16, "API", "인증 API", "로그인", "-", "-", "/api/auth/login", "R", "API_AUTH_002", "본인", "API"],
    [17, "API", "인증 API", "비밀번호 재설정 토큰 발급", "-", "-", "/api/auth/password/reset-request", "C", "API_AUTH_003", "본인", "API"],
    [18, "API", "인증 API", "비밀번호 재설정 확정", "-", "-", "/api/auth/password/reset-confirm", "R, U", "API_AUTH_004", "본인", "API"],
    [19, "API", "사용자 API", "내 정보 조회", "-", "-", "/api/me", "R", "API_USER_001", "본인", "API"],
    [20, "API", "사용자 API", "기본 주소 수정", "-", "-", "/api/me/address", "U", "API_USER_002", "본인", "API"],
    [21, "API", "친구 API", "친구 목록 조회", "-", "-", "/api/friends GET", "R", "API_FRIEND_001", "본인", "API"],
    [22, "API", "친구 API", "친구 추가", "-", "-", "/api/friends POST", "C", "API_FRIEND_002", "본인", "API"],
    [23, "API", "이력 API", "검색 이력 조회", "-", "-", "/api/history", "R", "API_HISTORY_001", "본인", "API"],
    [24, "API", "모임 API", "모임 목록 조회", "-", "-", "/api/meetings GET", "R", "API_MEETING_001", "본인", "API"],
    [25, "API", "모임 API", "모임 생성", "-", "-", "/api/meetings POST", "C", "API_MEETING_002", "본인", "API"],
    [26, "API", "추천 API", "카테고리 조회", "-", "-", "/api/categories", "R", "API_RECOMMEND_001", "본인", "API"],
    [27, "API", "추천 API", "장소 추천 조회", "-", "-", "/api/recommendations", "R", "API_RECOMMEND_002", "본인", "API"],
    [28, "API", "운영 API", "헬스체크", "-", "-", "/api/health", "-", "API_COMMON_001", "본인", "API"],
]

COLLECTION_HEADERS = ["Collection ID", "Collection 명", "NO", "Depth", "칼럼ID", "칼럼명", "Type", "Length", "Decimal", "PK", "NOT NULL", "UNIQUE", "비고"]
COLLECTION_ROWS = [
    ["REDIS_SESSION_STORE", "세션 토큰 저장소", 1, 1, "KEY", "세션 Redis Key (wemeet:session:{token})", "STRING", 255, "-", "*", "*", "", "token 기반 단건 조회"],
    ["", "", 2, 1, "VALUE_USER_ID", "저장 값(사용자 ID)", "STRING", 40, "-", "", "*", "", "Redis value 자체는 userId 문자열"],
    ["", "", 3, 1, "TTL_MINUTES", "만료시간(분)", "NUMBER", 10, 0, "", "*", "", "app.redis.session-ttl-minutes 설정값 사용"],
    ["REDIS_RECOMMENDATION_CACHE", "추천 결과 캐시", 1, 1, "KEY", "추천 Redis Key (wemeet:recommendation:{cacheKey})", "STRING", 255, "-", "*", "*", "", "cacheKey 기반 단건 조회"],
    ["", "", 2, 1, "VALUE_PAYLOAD_JSON", "저장 값(RecommendationResponse JSON)", "STRING", "4000+", "-", "", "*", "", "추천 응답 전체를 JSON 직렬화"],
    ["", "", 3, 1, "TTL_MINUTES", "만료시간(분)", "NUMBER", 10, 0, "", "*", "", "app.redis.recommendation-ttl-minutes 설정값 사용"],
]

TABLE_HEADERS = ["Table ID", "Table 명", "NO", "Column ID", "Column Name", "Type", "Length", "Decimal", "PK", "NOT NULL", "UNIQUE", "FK", "비고"]
IMPLEMENTED_TABLE_ROWS = [
    ["APP_USER", "회원정보", 1, "ID", "사용자 ID", "VARCHAR", 40, "-", "*", "*", "", "", ""],
    ["", "", 2, "LOGIN_ID", "로그인 ID", "VARCHAR", 60, "-", "", "*", "*", "", ""],
    ["", "", 3, "NICKNAME", "닉네임", "VARCHAR", 80, "-", "", "*", "", "", ""],
    ["", "", 4, "EMAIL", "이메일", "VARCHAR", 120, "-", "", "*", "*", "", ""],
    ["", "", 5, "PASSWORD_HASH", "비밀번호 해시", "VARCHAR", 120, "-", "", "*", "", "", ""],
    ["", "", 6, "FRIEND_CODE", "친구코드", "VARCHAR", 40, "-", "", "*", "*", "", ""],
    ["", "", 7, "BASE_ADDRESS", "기본 출발지 주소", "VARCHAR", 255, "-", "", "*", "", "", ""],
    ["", "", 8, "CREATED_AT", "생성일시", "DATETIME", "-", "-", "", "*", "", "", ""],
    ["", "", 9, "UPDATED_AT", "수정일시", "DATETIME", "-", "-", "", "*", "", "", ""],
    ["PASSWORD_RESET_TOKEN", "비밀번호 재설정 토큰", 1, "ID", "토큰 ID", "BIGINT", "-", "-", "*", "*", "", "", ""],
    ["", "", 2, "USER_ID", "사용자 ID", "VARCHAR", 40, "-", "", "*", "", "*", "APP_USER.ID"],
    ["", "", 3, "TOKEN_HASH", "토큰 해시", "VARCHAR", 80, "-", "", "*", "*", "", ""],
    ["", "", 4, "EXPIRES_AT", "만료일시", "DATETIME", "-", "-", "", "*", "", "", ""],
    ["", "", 5, "USED", "사용여부", "BOOLEAN", 1, "-", "", "*", "", "", ""],
    ["", "", 6, "CREATED_AT", "생성일시", "DATETIME", "-", "-", "", "*", "", "", ""],
    ["FRIEND_RELATION", "친구관계", 1, "USER_ID", "사용자 ID", "VARCHAR", 40, "-", "*", "*", "", "*", "APP_USER.ID"],
    ["", "", 2, "FRIEND_USER_ID", "친구 사용자 ID", "VARCHAR", 40, "-", "*", "*", "", "*", "APP_USER.ID"],
    ["", "", 3, "CREATED_AT", "생성일시", "DATETIME", "-", "-", "", "*", "", "", ""],
    ["SEARCH_HISTORY", "검색이력", 1, "HISTORY_ID", "검색이력 ID", "BIGINT", "-", "-", "*", "*", "", "", ""],
    ["", "", 2, "USER_ID", "사용자 ID", "VARCHAR", 40, "-", "", "*", "", "*", "APP_USER.ID"],
    ["", "", 3, "QUERY", "검색어", "VARCHAR", 255, "-", "", "*", "", "", ""],
    ["", "", 4, "CATEGORY", "카테고리", "VARCHAR", 30, "-", "", "*", "", "", ""],
    ["", "", 5, "SEARCHED_AT", "검색일시", "DATETIME", "-", "-", "", "*", "", "", ""],
    ["MEETING", "모임", 1, "MEETING_ID", "모임 ID", "VARCHAR", 40, "-", "*", "*", "", "", ""],
    ["", "", 2, "TITLE", "모임명", "VARCHAR", 120, "-", "", "*", "", "", ""],
    ["", "", 3, "DESCRIPTION", "설명", "VARCHAR", 1000, "-", "", "", "", "", ""],
    ["", "", 4, "MEETING_DATE", "모임일자", "DATE", "-", "-", "", "*", "", "", ""],
    ["", "", 5, "CATEGORY", "카테고리", "VARCHAR", 30, "-", "", "*", "", "", ""],
    ["", "", 6, "HOST_USER_ID", "주최자 ID", "VARCHAR", 40, "-", "", "*", "", "*", "APP_USER.ID"],
    ["", "", 7, "CREATED_AT", "생성일시", "DATETIME", "-", "-", "", "*", "", "", ""],
    ["MEETING_PARTICIPANT", "모임참가자", 1, "MEETING_ID", "모임 ID", "VARCHAR", 40, "-", "*", "*", "", "*", "MEETING.MEETING_ID"],
    ["", "", 2, "USER_ID", "사용자 ID", "VARCHAR", 40, "-", "*", "*", "", "*", "APP_USER.ID"],
    ["", "", 3, "ROLE", "역할(HOST/PARTICIPANT)", "VARCHAR", 20, "-", "", "*", "", "", ""],
    ["", "", 4, "CREATED_AT", "생성일시", "DATETIME", "-", "-", "", "*", "", "", ""],
]

LOGICAL_ENTITY_HEADERS = ["엔터티명", "설명"]
LOGICAL_ENTITY_ROWS = [
    ["APP_USER", "서비스 회원 기본 정보"],
    ["PASSWORD_RESET_TOKEN", "비밀번호 재설정 토큰 이력"],
    ["FRIEND_RELATION", "사용자 간 친구 관계"],
    ["SEARCH_HISTORY", "사용자 검색 이력"],
    ["MEETING", "모임 기본 정보"],
    ["MEETING_PARTICIPANT", "모임 참가자 관계"],
]
LOGICAL_REL_HEADERS = ["관계", "설명"]
LOGICAL_REL_ROWS = [
    ["APP_USER 1 : N PASSWORD_RESET_TOKEN", "한 사용자는 여러 개의 재설정 토큰을 가질 수 있다."],
    ["APP_USER N : N APP_USER", "친구 관계는 FRIEND_RELATION으로 해소한다."],
    ["APP_USER 1 : N SEARCH_HISTORY", "사용자별 검색 이력 저장."],
    ["APP_USER 1 : N MEETING", "한 사용자는 여러 모임을 생성할 수 있다."],
    ["APP_USER N : N MEETING", "참가 관계는 MEETING_PARTICIPANT로 해소한다."],
    ["MEETING 1 : N MEETING_PARTICIPANT", "한 모임에는 여러 참가자가 연결된다."],
]

PHYSICAL_KEY_HEADERS = ["구분", "컬럼", "비고"]
PHYSICAL_KEY_ROWS = [
    ["PK", "APP_USER.ID", "회원 기본키"],
    ["PK", "PASSWORD_RESET_TOKEN.ID", "토큰 기본키"],
    ["FK", "PASSWORD_RESET_TOKEN.USER_ID -> APP_USER.ID", "회원 참조"],
    ["FK", "FRIEND_RELATION.USER_ID -> APP_USER.ID", "친구관계 사용자 참조"],
    ["FK", "FRIEND_RELATION.FRIEND_USER_ID -> APP_USER.ID", "친구관계 친구 참조"],
    ["PK", "SEARCH_HISTORY.HISTORY_ID", "검색이력 기본키"],
    ["FK", "SEARCH_HISTORY.USER_ID -> APP_USER.ID", "회원 참조"],
    ["PK", "MEETING.MEETING_ID", "모임 기본키"],
    ["FK", "MEETING.HOST_USER_ID -> APP_USER.ID", "주최자 참조"],
    ["FK", "MEETING_PARTICIPANT.MEETING_ID -> MEETING.MEETING_ID", "모임 참조"],
    ["FK", "MEETING_PARTICIPANT.USER_ID -> APP_USER.ID", "참가자 참조"],
]
PHYSICAL_INDEX_HEADERS = ["인덱스명", "대상", "설명"]
PHYSICAL_INDEX_ROWS = [
    ["UK_APP_USER_LOGIN_ID", "APP_USER.LOGIN_ID", "로그인 ID 유니크"],
    ["UK_APP_USER_EMAIL", "APP_USER.EMAIL", "이메일 유니크"],
    ["UK_APP_USER_FRIEND_CODE", "APP_USER.FRIEND_CODE", "친구코드 유니크"],
    ["UK_PASSWORD_RESET_TOKEN_HASH", "PASSWORD_RESET_TOKEN.TOKEN_HASH", "토큰 해시 유니크"],
    ["IX_PASSWORD_RESET_TOKEN_USER_ID", "PASSWORD_RESET_TOKEN.USER_ID", "사용자별 재설정 토큰 조회"],
    ["UK_FRIEND_RELATION_USER_FRIEND", "FRIEND_RELATION(USER_ID, FRIEND_USER_ID)", "친구관계 중복 방지"],
    ["IX_FRIEND_RELATION_FRIEND_USER_ID", "FRIEND_RELATION.FRIEND_USER_ID", "반대 방향 친구 조회"],
    ["IX_SEARCH_HISTORY_USER_SEARCHED", "SEARCH_HISTORY(USER_ID, SEARCHED_AT)", "최근 검색 이력 조회"],
    ["IX_MEETING_HOST_USER_ID", "MEETING.HOST_USER_ID", "주최자별 모임 조회"],
    ["UK_MEETING_PARTICIPANT_MEETING_USER", "MEETING_PARTICIPANT(MEETING_ID, USER_ID)", "참가 중복 방지"],
    ["IX_MEETING_PARTICIPANT_USER_ID", "MEETING_PARTICIPANT.USER_ID", "사용자 참여 모임 조회"],
]

WBS_HEADERS = ["태스크", "작업자", "상태", "시작일", "종료일", "기간", "진척률"]
WBS_ROWS = [
    ["요구사항 정리", "본인", "Finished", date(2026, 3, 10), date(2026, 3, 11), 2, 1.00],
    ["화면 구조 설계", "본인", "Finished", date(2026, 3, 11), date(2026, 3, 12), 2, 1.00],
    ["메뉴 구조도 작성", "본인", "Finished", date(2026, 3, 11), date(2026, 3, 12), 2, 1.00],
    ["프로그램 명세서 작성", "본인", "Finished", date(2026, 3, 11), date(2026, 3, 12), 2, 1.00],
    ["DB 설계", "본인", "Finished", date(2026, 3, 12), date(2026, 3, 20), 9, 1.00],
    ["논리 ERD 작성", "본인", "Finished", date(2026, 3, 12), date(2026, 3, 20), 9, 1.00],
    ["물리 ERD 작성", "본인", "Finished", date(2026, 3, 13), date(2026, 3, 20), 8, 1.00],
    ["테이블 명세서 작성", "본인", "Finished", date(2026, 3, 13), date(2026, 3, 20), 8, 1.00],
    ["개발환경 세팅", "본인", "Finished", date(2026, 3, 12), date(2026, 3, 13), 2, 1.00],
    ["Spring Boot 세팅", "본인", "Finished", date(2026, 3, 12), date(2026, 3, 12), 1, 1.00],
    ["MariaDB 세팅", "본인", "Finished", date(2026, 3, 12), date(2026, 3, 13), 2, 1.00],
    ["Redis 세팅", "본인", "Finished", date(2026, 3, 12), date(2026, 3, 13), 2, 1.00],
    ["화면 구현", "본인", "Finished", date(2026, 3, 13), date(2026, 3, 20), 8, 1.00],
    ["로그인/회원가입 구현", "본인", "Finished", date(2026, 3, 13), date(2026, 3, 14), 2, 1.00],
    ["홈/친구/프로필 구현", "본인", "Finished", date(2026, 3, 14), date(2026, 3, 15), 2, 1.00],
    ["추천 결과/모임 구현", "본인", "Finished", date(2026, 3, 15), date(2026, 3, 20), 6, 1.00],
    ["API 구현", "본인", "Finished", date(2026, 3, 13), date(2026, 3, 20), 8, 1.00],
    ["인증 API 구현", "본인", "Finished", date(2026, 3, 13), date(2026, 3, 14), 2, 1.00],
    ["사용자/친구 API 구현", "본인", "Finished", date(2026, 3, 14), date(2026, 3, 15), 2, 1.00],
    ["모임/추천 API 구현", "본인", "Finished", date(2026, 3, 15), date(2026, 3, 17), 3, 1.00],
    ["배포 구성 정리", "본인", "Finished", date(2026, 3, 13), date(2026, 3, 14), 2, 1.00],
    ["Docker 구성", "본인", "Finished", date(2026, 3, 13), date(2026, 3, 13), 1, 1.00],
    ["AWS 배포 문서 정리", "본인", "Finished", date(2026, 3, 13), date(2026, 3, 14), 2, 1.00],
    ["테스트 코드 작성", "본인", "In Progress", date(2026, 3, 14), date(2026, 3, 20), 7, 0.75],
    ["문서 보정 및 제출본 정리", "본인", "Finished", date(2026, 3, 17), date(2026, 3, 20), 4, 1.00],
]


def style_title(ws, row: int, text: str, width: int) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    cell = ws.cell(row=row, column=1, value=text)
    cell.fill = TITLE_FILL
    cell.font = Font(color="FFFFFF", bold=True, size=14)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = THIN_BORDER
    ws.row_dimensions[row].height = 24


def style_table(ws, headers: list[str], rows: list[list], start_row: int) -> int:
    for col_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col_index, value=header)
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    for row_index, row in enumerate(rows, start=start_row + 1):
        for col_index, value in enumerate(row, start=1):
            cell = ws.cell(row=row_index, column=col_index, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN_BORDER
    return start_row + len(rows)


def add_note(ws, row: int, text: str, width: int) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    cell = ws.cell(row=row, column=1, value=text)
    cell.fill = SUBSECTION_FILL
    cell.font = Font(italic=True, color="44546A")
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    cell.border = THIN_BORDER


def set_widths(ws, widths: list[int]) -> None:
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def create_menu_sheet(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "메뉴구조도"
    style_title(ws, 1, "WeMeet 개발산출물 - 메뉴구조도", len(MENU_HEADERS))
    add_note(ws, 2, "현재 구현 기준 화면/탭/API 흐름을 정리한 메뉴 구조도", len(MENU_HEADERS))
    end_row = style_table(ws, MENU_HEADERS, MENU_ROWS, 4)
    ws.freeze_panes = "A5"
    set_widths(ws, [12, 18, 18, 20, 12, 12, 12, 14, 30])
    ws.auto_filter.ref = f"A4:I{end_row}"


def create_program_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("프로그램 명세서")
    style_title(ws, 1, "WeMeet 개발산출물 - 프로그램 명세서", len(PROGRAM_HEADERS))
    add_note(ws, 2, "화면 프로그램과 REST API를 함께 정리한 프로그램 명세서", len(PROGRAM_HEADERS))
    end_row = style_table(ws, PROGRAM_HEADERS, PROGRAM_ROWS, 4)
    ws.freeze_panes = "A5"
    set_widths(ws, [8, 18, 18, 18, 14, 14, 28, 16, 18, 10, 10])
    ws.auto_filter.ref = f"A4:K{end_row}"


def create_collection_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("컬렉션 정의서(NoSQL)")
    style_title(ws, 1, "WeMeet 개발산출물 - 컬렉션 정의서(NoSQL)", len(COLLECTION_HEADERS))
    add_note(ws, 2, "MongoDB 대신 Redis 세션/추천 캐시 키 구조를 정리한 대체 시트", len(COLLECTION_HEADERS))
    end_row = style_table(ws, COLLECTION_HEADERS, COLLECTION_ROWS, 4)
    ws.freeze_panes = "A5"
    set_widths(ws, [22, 18, 8, 8, 22, 30, 12, 10, 10, 8, 10, 10, 18])
    ws.auto_filter.ref = f"A4:M{end_row}"


def create_table_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("테이블 명세서(RDBMS)")
    style_title(ws, 1, "WeMeet 개발산출물 - 테이블 명세서(RDBMS)", len(TABLE_HEADERS))
    add_note(ws, 2, "현재 MariaDB/JPA 기준 실제 생성 테이블을 정리", len(TABLE_HEADERS))

    ws.merge_cells("A4:M4")
    ws["A4"] = "4.1 구현 테이블"
    ws["A4"].fill = SECTION_FILL
    ws["A4"].font = Font(bold=True)
    ws["A4"].alignment = Alignment(horizontal="left", vertical="center")
    ws["A4"].border = THIN_BORDER
    end_row = style_table(ws, TABLE_HEADERS, IMPLEMENTED_TABLE_ROWS, 5)

    ws.freeze_panes = "A6"
    set_widths(ws, [22, 18, 8, 24, 24, 14, 10, 10, 8, 10, 10, 18, 28])
    ws.auto_filter.ref = f"A5:M{end_row}"


def create_logical_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("논리 데이터 모델링(ERD)")
    style_title(ws, 1, "WeMeet 개발산출물 - 논리 데이터 모델링(ERD)", 4)
    add_note(ws, 2, "ERD 이미지 대신 현재 구현 기준 엔터티/관계 요약을 표로 정리", 4)

    ws.merge_cells("A4:B4")
    ws["A4"] = "엔터티 목록"
    ws["A4"].fill = SECTION_FILL
    ws["A4"].font = Font(bold=True)
    ws["A4"].border = THIN_BORDER
    style_table(ws, LOGICAL_ENTITY_HEADERS, LOGICAL_ENTITY_ROWS, 5)

    ws.merge_cells("A13:B13")
    ws["A13"] = "관계 목록"
    ws["A13"].fill = SECTION_FILL
    ws["A13"].font = Font(bold=True)
    ws["A13"].border = THIN_BORDER
    style_table(ws, LOGICAL_REL_HEADERS, LOGICAL_REL_ROWS, 14)
    set_widths(ws, [34, 66, 14, 14])


def create_physical_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("물리 데이터 모델링(ERD)")
    style_title(ws, 1, "WeMeet 개발산출물 - 물리 데이터 모델링(ERD)", 4)
    add_note(ws, 2, "PK/FK 구성과 인덱스 제안을 물리 모델 기준으로 요약", 4)

    ws.merge_cells("A4:C4")
    ws["A4"] = "PK/FK 요약"
    ws["A4"].fill = SECTION_FILL
    ws["A4"].font = Font(bold=True)
    ws["A4"].border = THIN_BORDER
    style_table(ws, PHYSICAL_KEY_HEADERS, PHYSICAL_KEY_ROWS, 5)

    ws.merge_cells("A18:C18")
    ws["A18"] = "인덱스 제안"
    ws["A18"].fill = SECTION_FILL
    ws["A18"].font = Font(bold=True)
    ws["A18"].border = THIN_BORDER
    style_table(ws, PHYSICAL_INDEX_HEADERS, PHYSICAL_INDEX_ROWS, 19)
    set_widths(ws, [16, 46, 34, 14])


def create_wbs_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("WBS")
    style_title(ws, 1, "WeMeet 개발산출물 - WBS", len(WBS_HEADERS))
    add_note(ws, 2, "2026-03-20 기준 구현 상태를 반영한 작업 일정/진척률", len(WBS_HEADERS))
    end_row = style_table(ws, WBS_HEADERS, WBS_ROWS, 4)
    for row in range(5, end_row + 1):
        ws.cell(row=row, column=4).number_format = "yyyy-mm-dd"
        ws.cell(row=row, column=5).number_format = "yyyy-mm-dd"
        ws.cell(row=row, column=7).number_format = "0%"
    ws.freeze_panes = "A5"
    set_widths(ws, [28, 12, 14, 14, 14, 10, 10])
    ws.auto_filter.ref = f"A4:G{end_row}"


def build_workbook() -> Workbook:
    wb = Workbook()
    create_menu_sheet(wb)
    create_program_sheet(wb)
    create_collection_sheet(wb)
    create_table_sheet(wb)
    create_logical_sheet(wb)
    create_physical_sheet(wb)
    create_wbs_sheet(wb)
    return wb


def main() -> None:
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb = build_workbook()
    wb.save(OUTPUT_PATH)
    print(f"Saved workbook to {OUTPUT_PATH.resolve()}")


if __name__ == "__main__":
    main()
