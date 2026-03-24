from __future__ import annotations

from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import generate_deliverables as gd


TEMPLATE_PATH = Path("artifacts/2420110158_김소연_개발산출물.xlsx")
OUTPUT_PATH = Path("artifacts/2420110158_김소연_개발산출물_wemeet.xlsx")
WORKER_NAME = "김현규"

HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
SECTION_FILL = PatternFill("solid", fgColor="EAF2F8")
THIN_BORDER = Border(
    left=Side(style="thin", color="B7C3D0"),
    right=Side(style="thin", color="B7C3D0"),
    top=Side(style="thin", color="B7C3D0"),
    bottom=Side(style="thin", color="B7C3D0"),
)


def copy_row_style(ws, source_row: int, target_row: int, start_col: int, end_col: int) -> None:
    for col in range(start_col, end_col + 1):
        source = ws.cell(row=source_row, column=col)
        target = ws.cell(row=target_row, column=col)
        target._style = copy(source._style)
        if source.has_style and source.alignment:
            target.alignment = copy(source.alignment)
        if source.font:
            target.font = copy(source.font)
        if source.fill:
            target.fill = copy(source.fill)
        if source.border:
            target.border = copy(source.border)
        if source.number_format:
            target.number_format = source.number_format
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def ensure_capacity(ws, required_max_row: int, style_source_row: int, start_col: int, end_col: int) -> None:
    current_max_row = ws.max_row
    if current_max_row >= required_max_row:
        return
    ws.insert_rows(current_max_row + 1, required_max_row - current_max_row)
    for row in range(current_max_row + 1, required_max_row + 1):
        copy_row_style(ws, style_source_row, row, start_col, end_col)


def clear_values(ws, start_row: int, end_row: int, start_col: int, end_col: int) -> None:
    for row in range(start_row, end_row + 1):
        for col in range(start_col, end_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.value = None
            cell.hyperlink = None
            cell.comment = None


def unmerge_from_row(ws, start_row: int) -> None:
    for merged_range in list(ws.merged_cells.ranges):
        if merged_range.min_row >= start_row:
            ws.unmerge_cells(str(merged_range))


def merge_repeated_values(ws, start_row: int, end_row: int, columns: list[int]) -> None:
    for col in columns:
        span_start = start_row
        previous = ws.cell(row=start_row, column=col).value
        for row in range(start_row + 1, end_row + 2):
            current = ws.cell(row=row, column=col).value if row <= end_row else object()
            if current != previous or previous in (None, "", "-"):
                if previous not in (None, "", "-") and row - 1 > span_start:
                    ws.merge_cells(
                        start_row=span_start,
                        start_column=col,
                        end_row=row - 1,
                        end_column=col,
                    )
                    ws.cell(row=span_start, column=col).alignment = Alignment(
                        horizontal="center",
                        vertical="center",
                        wrap_text=True,
                    )
                span_start = row
                previous = current


def write_offset_rows(
    ws,
    rows: list[list],
    start_row: int,
    start_col: int,
    style_source_row: int,
    merge_columns: list[int] | None = None,
) -> int:
    end_row = start_row + len(rows) - 1
    end_col = start_col + len(rows[0]) - 1
    ensure_capacity(ws, end_row, style_source_row, start_col, end_col)
    unmerge_from_row(ws, start_row)
    clear_values(ws, start_row, ws.max_row, start_col, end_col)

    for row_index, row_values in enumerate(rows, start=start_row):
        for col_index, value in enumerate(row_values, start=start_col):
            ws.cell(row=row_index, column=col_index, value=value)

    if merge_columns:
        merge_repeated_values(ws, start_row, end_row, [start_col + offset for offset in merge_columns])

    return end_row


def workerized_program_rows(worker_name: str) -> list[list]:
    rows = []
    for row in gd.PROGRAM_ROWS:
        values = list(row)
        values[9] = worker_name
        rows.append(values)
    return rows


def workerized_wbs_rows(worker_name: str) -> list[list]:
    rows = []
    for row in gd.WBS_ROWS:
        values = list(row)
        values[1] = worker_name
        rows.append(values)
    return rows


def fill_menu_sheet(workbook) -> None:
    ws = workbook["메뉴구조도"]
    ws["B2"] = "WeMeet 사용자/API"
    menu_rows = [list(row) for row in gd.MENU_ROWS]
    end_row = write_offset_rows(ws, menu_rows, start_row=5, start_col=2, style_source_row=5, merge_columns=[0, 1, 2, 3])
    ws.auto_filter.ref = f"B4:J{end_row}"


def fill_program_sheet(workbook, worker_name: str) -> None:
    ws = workbook["프로그램 명세서"]
    ws["B2"] = "WeMeet 사용자/API"
    program_rows = workerized_program_rows(worker_name)
    end_row = write_offset_rows(ws, program_rows, start_row=5, start_col=2, style_source_row=5, merge_columns=[1, 2, 3, 4, 5])
    ws.auto_filter.ref = f"B3:L{end_row}"


def fill_collection_sheet(workbook) -> None:
    ws = workbook["컬렉션 정의서(NoSQL)"]
    ws["B2"] = "컬랙션 정의서"
    collection_rows = [list(row[:-1]) for row in gd.COLLECTION_ROWS]
    end_row = write_offset_rows(ws, collection_rows, start_row=4, start_col=2, style_source_row=4, merge_columns=[0, 1])
    ws.auto_filter.ref = f"B3:M{end_row}"


def fill_table_sheet(workbook) -> None:
    ws = workbook["테이블 명세서(RDBMS)"]
    ws["B2"] = "테이블 명세서"
    table_rows = [list(row) for row in gd.IMPLEMENTED_TABLE_ROWS]
    end_row = write_offset_rows(ws, table_rows, start_row=4, start_col=2, style_source_row=4, merge_columns=[0, 1])
    ws.auto_filter.ref = f"B3:N{end_row}"


def write_section_title(ws, row: int, title: str, start_col: int, end_col: int) -> None:
    ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
    cell = ws.cell(row=row, column=start_col, value=title)
    cell.fill = SECTION_FILL
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    cell.border = THIN_BORDER


def write_simple_table(ws, start_row: int, start_col: int, headers: list[str], rows: list[list]) -> int:
    header_row = start_row
    for offset, header in enumerate(headers):
        cell = ws.cell(row=header_row, column=start_col + offset, value=header)
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    for row_index, row_values in enumerate(rows, start=header_row + 1):
        for offset, value in enumerate(row_values):
            cell = ws.cell(row=row_index, column=start_col + offset, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN_BORDER
    return header_row + len(rows)


def fill_logical_sheet(workbook) -> None:
    ws = workbook["논리 데이터 모델링(ERD)"]
    ws.delete_rows(1, ws.max_row)
    write_section_title(ws, 2, "엔터티 목록", 2, 3)
    entity_end = write_simple_table(ws, 3, 2, gd.LOGICAL_ENTITY_HEADERS, [list(row) for row in gd.LOGICAL_ENTITY_ROWS])
    write_section_title(ws, entity_end + 2, "관계 목록", 2, 3)
    write_simple_table(ws, entity_end + 3, 2, gd.LOGICAL_REL_HEADERS, [list(row) for row in gd.LOGICAL_REL_ROWS])
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 72


def fill_physical_sheet(workbook) -> None:
    ws = workbook["물리 데이터 모델링(ERD)"]
    ws.delete_rows(1, ws.max_row)
    write_section_title(ws, 2, "PK/FK 요약", 2, 4)
    key_end = write_simple_table(ws, 3, 2, gd.PHYSICAL_KEY_HEADERS, [list(row) for row in gd.PHYSICAL_KEY_ROWS])
    write_section_title(ws, key_end + 2, "인덱스 제안", 2, 4)
    write_simple_table(ws, key_end + 3, 2, gd.PHYSICAL_INDEX_HEADERS, [list(row) for row in gd.PHYSICAL_INDEX_ROWS])
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 52
    ws.column_dimensions["D"].width = 36


def fill_wbs_sheet(workbook, worker_name: str) -> None:
    ws = workbook["WBS"]
    wbs_rows = workerized_wbs_rows(worker_name)
    required_max_row = 1 + len(wbs_rows)
    ensure_capacity(ws, required_max_row, 3, 1, 7)
    for row in range(2, ws.max_row + 1):
        copy_row_style(ws, 3, row, 1, 7)
    clear_values(ws, 2, ws.max_row, 1, 7)

    for row_index, row_values in enumerate(wbs_rows, start=2):
        for col_index, value in enumerate(row_values, start=1):
            cell = ws.cell(row=row_index, column=col_index, value=value)
            if col_index in (4, 5):
                cell.number_format = "yyyy-mm-dd"
            if col_index == 7:
                cell.number_format = "0%"


def main() -> None:
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"Template not found: {TEMPLATE_PATH}")

    workbook = load_workbook(TEMPLATE_PATH)
    fill_menu_sheet(workbook)
    fill_program_sheet(workbook, WORKER_NAME)
    fill_collection_sheet(workbook)
    fill_table_sheet(workbook)
    fill_logical_sheet(workbook)
    fill_physical_sheet(workbook)
    fill_wbs_sheet(workbook, WORKER_NAME)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT_PATH)
    print(f"Saved workbook to {OUTPUT_PATH.resolve()}")


if __name__ == "__main__":
    main()
